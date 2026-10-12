import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('d:/BTL_CumTramCiTiBike/LuuVanDucAnh_HocMayCoBan.docx')

print('=== GENERAL STATS ===')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')

print('\n=== ALL HEADINGS (C1, C2, C3) ===')
for i, p in enumerate(doc.paragraphs):
    if p.style.name in ['C1', 'C2', 'C3']:
        ind = p._p.xpath('w:pPr/w:ind')
        ind_str = ind[0].attrib if ind else {}
        left_val = ind_str.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}left', '0')
        print(f'{i:3d} | {p.style.name:3s} | ind_left={left_val} | {p.text}')

print('\n=== ALL TABLES (0 to 4) ===')
for i, tbl in enumerate(doc.tables):
    prev_p = tbl._element.getprevious()
    prev_txt = docx.text.paragraph.Paragraph(prev_p, doc).text if prev_p is not None else 'NONE'
    c0 = tbl.rows[0].cells[0]
    shd = c0._tc.xpath('w:tcPr/w:shd')
    shd_fill = shd[0].attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd else 'none'
    borders = tbl._element.xpath('w:tblPr/w:tblBorders')
    has_borders = len(borders) > 0
    print(f'Table {i}: rows={len(tbl.rows):2d}, cols={len(tbl.columns)}, align={tbl.alignment}, header_fill={shd_fill}, borders={has_borders}')
    print(f'   Caption ABOVE: {repr(prev_txt)}')

print('\n=== ALL FIGURES (Images + Captions) ===')
img_idx = 0
for i, p in enumerate(doc.paragraphs):
    blips = p._p.xpath('.//a:blip')
    if blips and not any(r.text for r in p.runs):
        img_idx += 1
        exts = p._p.xpath('.//wp:extent')
        cx = int(exts[0].attrib['cx']) / 914400 if exts else 0
        cy = int(exts[0].attrib['cy']) / 914400 if exts else 0
        next_p = doc.paragraphs[i+1] if i+1 < len(doc.paragraphs) else None
        next_txt = next_p.text if next_p else 'NONE'
        next_style = next_p.style.name if next_p else 'NONE'
        print(f'Figure {img_idx} at P {i:3d} [width={cx:.2f}in, height={cy:.2f}in, align={p.alignment}]:')
        print(f'   Caption BELOW [{next_style}]: {next_txt}')

print('\n=== FRONT MATTER LISTS ===')
for i in range(70, 95):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f'{i:2d} [{p.style.name:16s}]: {p.text}')
