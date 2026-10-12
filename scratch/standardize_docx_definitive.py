import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

SRC_DOCX = "d:/BTL_CumTramCiTiBike/LuuVanDucAnh_HocMayCoBan_pre_style_fix.docx"
OUT_DOCX_PRIMARY = "d:/BTL_CumTramCiTiBike/LuuVanDucAnh_HocMayCoBan.docx"
OUT_DOCX_BACKUP = "d:/BTL_CumTramCiTiBike/LuuVanDucAnh_HocMayCoBan_ChuanHoa.docx"
FIG_DIR = "d:/BTL_CumTramCiTiBike/reports/figures"

print("Loading clean baseline document from:", SRC_DOCX)
doc = docx.Document(SRC_DOCX)

# ==============================================================================
# Helper Formatting Functions
# ==============================================================================

def clean_paragraph_indent(p):
    """Remove any irregular hanging/left/right indents on heading paragraphs."""
    pPr = p._p.get_or_add_pPr()
    for ind in pPr.xpath('w:ind'):
        pPr.remove(ind)
    new_ind = OxmlElement('w:ind')
    new_ind.set(qn('w:left'), '0')
    new_ind.set(qn('w:right'), '0')
    new_ind.set(qn('w:firstLine'), '0')
    pPr.append(new_ind)

def set_paragraph_spacing(p, before_pt, after_pt, keep_with_next=True):
    """Set uniform spacing and keep-with-next property."""
    pPr = p._p.get_or_add_pPr()
    for sp in pPr.xpath('w:spacing'):
        pPr.remove(sp)
    new_sp = OxmlElement('w:spacing')
    new_sp.set(qn('w:before'), str(int(before_pt * 20)))
    new_sp.set(qn('w:after'), str(int(after_pt * 20)))
    new_sp.set(qn('w:line'), '280') # 1.15 - 1.2 line spacing
    new_sp.set(qn('w:lineRule'), 'auto')
    pPr.append(new_sp)
    
    if keep_with_next:
        for kwn in pPr.xpath('w:keepNext'):
            pPr.remove(kwn)
        pPr.append(OxmlElement('w:keepNext'))

def set_heading_text(p, text, font_size_pt=13, bold=True, color_rgb=None):
    """Format heading with exact typography."""
    p.text = ""
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb

def set_caption_text(p, prefix, title):
    """Format figure caption with bold prefix and italic title."""
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = "a hình"
    
    pPr = p._p.get_or_add_pPr()
    for sp in pPr.xpath('w:spacing'):
        pPr.remove(sp)
    new_sp = OxmlElement('w:spacing')
    new_sp.set(qn('w:before'), '40')  # 2pt
    new_sp.set(qn('w:after'), '160')  # 8pt
    new_sp.set(qn('w:line'), '240')
    new_sp.set(qn('w:lineRule'), 'auto')
    pPr.append(new_sp)
    
    # Run 1: Bold prefix
    r1 = p.add_run(prefix)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    # Run 2: Italic title
    r2 = p.add_run(" " + title)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.italic = True

def add_table_caption_above(tbl, caption_prefix, caption_title):
    """Insert standardized Table Caption paragraph above table."""
    p_elem = OxmlElement('w:p')
    tbl._element.addprevious(p_elem)
    p = docx.text.paragraph.Paragraph(p_elem, doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '160') # 8pt
    sp.set(qn('w:after'), '80')   # 4pt
    sp.set(qn('w:line'), '240')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    pPr.append(OxmlElement('w:keepNext'))
    
    r1 = p.add_run(caption_prefix + ": ")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r2 = p.add_run(caption_title)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.bold = True
    return p

def format_academic_table(tbl, col_widths, col_alignments):
    """Apply uniform APA academic styling to a table."""
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    tblPr = tbl._element.xpath('w:tblPr')
    if tblPr:
        for b in tblPr[0].xpath('w:tblBorders'):
            tblPr[0].remove(b)
        tblBorders = OxmlElement('w:tblBorders')
        
        for name, val, sz, col in [
            ('top', 'single', '10', '2A4B7C'),
            ('bottom', 'single', '10', '2A4B7C'),
            ('left', 'none', '0', 'auto'),
            ('right', 'none', '0', 'auto'),
            ('insideH', 'single', '4', 'CBD5E1'),
            ('insideV', 'none', '0', 'auto')
        ]:
            border_el = OxmlElement(f'w:{name}')
            border_el.set(qn('w:val'), val)
            border_el.set(qn('w:sz'), sz)
            border_el.set(qn('w:space'), '0')
            border_el.set(qn('w:color'), col)
            tblBorders.append(border_el)
            
        tblPr[0].append(tblBorders)
    
    for r_idx, row in enumerate(tbl.rows):
        trPr = row._tr.get_or_add_trPr()
        for cs in trPr.xpath('w:cantSplit'):
            trPr.remove(cs)
        trPr.append(OxmlElement('w:cantSplit'))
        
        if r_idx == 0:
            for th in trPr.xpath('w:tblHeader'):
                trPr.remove(th)
            trPr.append(OxmlElement('w:tblHeader'))
            
        for c_idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            
            if c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is not None:
                    tcW.set(qn('w:w'), str(int(col_widths[c_idx] * 1440)))
            
            for va in tcPr.xpath('w:vAlign'):
                tcPr.remove(va)
            vAlign = OxmlElement('w:vAlign')
            vAlign.set(qn('w:val'), 'center')
            tcPr.append(vAlign)
            
            for mar in tcPr.xpath('w:tcMar'):
                tcPr.remove(mar)
            tcMar = OxmlElement('w:tcMar')
            for m, val in [('top', '100'), ('bottom', '100'), ('left', '150'), ('right', '150')]:
                node = OxmlElement(f'w:{m}')
                node.set(qn('w:w'), val)
                node.set(qn('w:type'), 'dxa')
                tcMar.append(node)
            tcPr.append(tcMar)
            
            for s in tcPr.xpath('w:shd'):
                tcPr.remove(s)
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            if r_idx == 0:
                shd.set(qn('w:fill'), 'F0F4F8')
            else:
                shd.set(qn('w:fill'), 'FFFFFF')
            tcPr.append(shd)
            
            for p in cell.paragraphs:
                pPr = p._p.get_or_add_pPr()
                for sp in pPr.xpath('w:spacing'):
                    pPr.remove(sp)
                sp = OxmlElement('w:spacing')
                sp.set(qn('w:before'), '40')
                sp.set(qn('w:after'), '40')
                sp.set(qn('w:line'), '240')
                sp.set(qn('w:lineRule'), 'auto')
                pPr.append(sp)
                
                if r_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    align = col_alignments[c_idx] if c_idx < len(col_alignments) else WD_ALIGN_PARAGRAPH.LEFT
                    p.alignment = align
                    
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10.5)
                    if r_idx == 0:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

# ==============================================================================
# 1. STANDARDIZE HEADINGS (ONLY ON PARAGRAPHS WITH HEADING STYLES C1, C2, C3)
# ==============================================================================
print("Standardizing heading paragraphs...")

heading_map = {
    # C1
    ("C1", "CHƯƠNG 1: TỔNG QUAN VỀ ĐỀ TÀI"): ("CHƯƠNG 1: TỔNG QUAN VỀ ĐỀ TÀI", 14),
    ("C1", "CƠ SỞ LÝ THUYẾT"): ("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT", 14),
    ("C1", "XÂY DỰNG MÔ HÌNH HỌC MÁY VÀ 4 THÍ NGHIỆM BẮT BUỘC"): ("CHƯƠNG 3: XÂY DỰNG MÔ HÌNH HỌC MÁY VÀ 4 THÍ NGHIỆM BẮT BUỘC", 14),
    ("C1", "XÂY DỰNG ỨNG DỤNG WEB PHỤC VỤ DỰ ĐOÁN VÀ TRỰC QUAN HÓA"): ("CHƯƠNG 4: XÂY DỰNG ỨNG DỤNG WEB PHỤC VỤ DỰ ĐOÁN VÀ TRỰC QUAN HÓA", 14),
    ("C1", "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN"): ("CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", 14),
    ("C1", "TÀI LIỆU THAM KHẢO"): ("TÀI LIỆU THAM KHẢO", 14),

    # C2
    ("C2", "Lý do chọn đề tài"): ("1.1. Lý do chọn đề tài", 13),
    ("C2", "Mục tiêu của đề tài"): ("1.2. Mục tiêu của đề tài", 13),
    ("C2", "Nội dung thực hiện"): ("1.3. Nội dung thực hiện", 13),
    ("C2", "Phương pháp tiếp cận"): ("1.4. Phương pháp tiếp cận", 13),
    
    ("C2", "Khái niệm khai phá dữ liệu và học máy"): ("2.1. Khái niệm khai phá dữ liệu và học máy", 13),
    ("C2", "Thuật toán phân cụm K_Means"): ("2.2. Thuật toán phân cụm K-Means", 13),
    ("C2", "Thuật toán phân cụm K-Means"): ("2.2. Thuật toán phân cụm K-Means", 13),
    ("C2", "Các phương pháp lựa chọn số cụm K và đánh giá chất lượng phân cụm nội tại"): ("2.3. Các phương pháp lựa chọn số cụm K và đánh giá chất lượng phân cụm nội tại", 13),
    ("C2", "Đánh giá độ ổn định của cụm và các độ đo tương đồng ngoại tại"): ("2.4. Đánh giá độ ổn định của cụm và các độ đo tương đồng ngoại tại", 13),
    
    ("C2", "Thu thập dữ liệu và Kiểm tra toàn vẹn"): ("3.1. Thu thập dữ liệu và Kiểm tra toàn vẹn", 13),
    ("C2", "Chiến lược phân chia dữ liệu chống rò rỉ (Anti-Leakage)"): ("3.2. Chiến lược phân chia dữ liệu chống rò rỉ (Anti-Leakage)", 13),
    ("C2", "Trích xuất đặc trưng và Phân tích khám phá dữ liệu (EDA) trên tập Train"): ("3.3. Trích xuất đặc trưng và Phân tích khám phá dữ liệu (EDA)", 13),
    ("C2", "Bốn thí nghiệm khoa học bắt buộc"): ("3.4. Bốn thí nghiệm khoa học bắt buộc", 13),
    ("C2", "Kết quả mở niêm phong tập Test độc lập (N = 335 trạm)"): ("3.5. Kết quả mở niêm phong tập Test độc lập (N = 335 trạm)", 13),
    ("C2", "Đóng gói mô hình và Tiêu chuẩn hóa Model Card"): ("3.6. Đóng gói mô hình và Tiêu chuẩn hóa Model Card", 13),
    
    ("C2", "Kiến trúc hệ thống Serving"): ("4.1. Kiến trúc hệ thống Serving", 13),
    ("C2", "Phát triển Backend API (FastAPI)"): ("4.2. Phát triển Backend API (FastAPI)", 13),
    ("C2", "Phát triển Giao diện Web tương tác (3 Màn hình)"): ("4.3. Phát triển Giao diện Web tương tác (3 Màn hình)", 13),
    ("C2", "Kiểm thử tự động hệ thống (Testing)"): ("4.4. Kiểm thử tự động hệ thống (Testing)", 13),
    
    ("C2", "Kết quả đạt được"): ("5.1. Kết quả đạt được", 13),
    ("C2", "Hạn chế của đề tài"): ("5.2. Hạn chế của đề tài", 13),
    ("C2", "Hướng phát triển trong tương lai"): ("5.3. Hướng phát triển trong tương lai", 13),

    # C3
    ("C3", "Mục tiêu tổng quát"): ("1.2.1. Mục tiêu tổng quát", 13),
    ("C3", "Mục tiêu cụ thể"): ("1.2.2. Mục tiêu cụ thể", 13),
    ("C3", "Đối tượng và Phạm vi nghiên cứu"): ("1.2.3. Đối tượng và Phạm vi nghiên cứu", 13),
    
    ("C3", "Khái niệm khai phá dữ liệu"): ("2.1.1. Khái niệm khai phá dữ liệu trong giao thông đô thị", 13),
    ("C3", "Khái niệm học máy và bài toán phân loại"): ("2.1.2. Khái niệm học máy và bài toán phân cụm không giám sát", 13),
    ("C3", "Hàm mục tiêu và nguyên lý toán học"): ("2.2.1. Hàm mục tiêu và nguyên lý toán học", 13),
    ("C3", "Thuật toán tối ưu hóa Lloyd"): ("2.2.2. Thuật toán tối ưu hóa Lloyd", 13),
    ("C3", "Phương pháp khởi tạo K-Means++ và hạn chế của khởi tạo ngẫu nhiên"): ("2.2.3. Phương pháp khởi tạo K-Means++ và hạn chế của khởi tạo ngẫu nhiên", 13),
    ("C3", "Phương pháp điểm khuỷu tay Elbow"): ("2.3.1. Phương pháp điểm khuỷu tay (Elbow Method)", 13),
    ("C3", "Hệ số Silhouette"): ("2.3.2. Hệ số Silhouette (Silhouette Coefficient)", 13),
    ("C3", "Chỉ số Davies-Bouldin Index DBI"): ("2.3.3. Chỉ số Davies-Bouldin (DBI)", 13),
    ("C3", "Adjusted Rand Index ARI"): ("2.4.1. Chỉ số hiệu chỉnh Rand (Adjusted Rand Index - ARI)", 13),
    ("C3", "Normalized Mutual Information NMI"): ("2.4.2. Độ tương hỗ chuẩn hóa (Normalized Mutual Information - NMI)", 13),
    
    ("C3", "Thí nghiệm 1: So sánh Volume Baseline với Hồ sơ Tỷ lệ"): ("3.4.1. Thí nghiệm 1: So sánh mô hình cơ sở Volume Baseline với mô hình Profile K-Means", 13),
    ("C3", "Thí nghiệm 2: Khảo sát đường cong Elbow và Silhouette Score (K = 2 đến 8)"): ("3.4.2. Thí nghiệm 2: Khảo sát đường cong Elbow và Silhouette Score (K = 2 đến 8)", 13),
    ("C3", "Thí nghiệm 3: Khảo sát độ ổn định qua Seed và Chuyển dịch theo mùa"): ("3.4.3. Thí nghiệm 3: Khảo sát độ ổn định qua Seed và Chuyển dịch theo mùa", 13),
    ("C3", "Thí nghiệm 4: Mô tả cụm bằng bảng Median và Bản đồ trực quan"): ("3.4.4. Thí nghiệm 4: Mô tả cụm bằng bảng Median và Bản đồ trực quan", 13)
}

for p in doc.paragraphs:
    if p.style.name in ["C1", "C2", "C3"]:
        raw = p.text.strip()
        matched = False
        for (st_key, title_key), (new_title, font_size) in heading_map.items():
            if p.style.name == st_key and (raw == title_key or raw.startswith(title_key)):
                clean_paragraph_indent(p)
                if st_key == "C1":
                    set_paragraph_spacing(p, before_pt=18, after_pt=6, keep_with_next=True)
                    color = RGBColor(0x1E, 0x3A, 0x8A)
                elif st_key == "C2":
                    set_paragraph_spacing(p, before_pt=12, after_pt=4, keep_with_next=True)
                    color = RGBColor(0x1F, 0x29, 0x37)
                else:
                    set_paragraph_spacing(p, before_pt=8, after_pt=3, keep_with_next=True)
                    color = RGBColor(0x37, 0x41, 0x51)
                set_heading_text(p, new_title, font_size_pt=font_size, bold=True, color_rgb=color)
                matched = True
                break
        if not matched:
            # Still clean indentation and spacing for any custom heading
            clean_paragraph_indent(p)
            set_paragraph_spacing(p, before_pt=8, after_pt=3, keep_with_next=True)

# Clean up typos in body paragraphs
for p in doc.paragraphs:
    if "Đối Đối tượng nghiên cứu:" in p.text:
        p.text = p.text.replace("Đối Đối tượng nghiên cứu:", "Đối tượng nghiên cứu:")
    if "Trong y sinh học, khai phá dữ liệu giúp tìm kiếm các đặc trưng hình học" in p.text:
        p.text = ("Khai phá dữ liệu (Data Mining) là quá trình tự động hoặc bán tự động trích xuất các tri thức, "
                  "quy luật tiềm ẩn (patterns), các mối tương quan có ý nghĩa và thông tin hữu ích từ các tập dữ liệu lớn. "
                  "Trong lĩnh vực giao thông đô thị thông minh và hệ thống chia sẻ xe đạp công cộng (Bike-sharing Systems), "
                  "khai phá dữ liệu giúp nhận diện nhịp sống đô thị, thói quen di chuyển và các đặc trưng nhu cầu sử dụng "
                  "tại từng trạm xe đạp từ hàng triệu bản ghi hành trình thực tế.")
    if "Đề tài thuộc bài toán Học có giám sát (Supervised Learning) dạng Phân loại nhị phân" in p.text:
        p.text = ("Học máy (Machine Learning) là một phân nhánh cốt lõi của Trí tuệ nhân tạo, tập trung vào việc nghiên "
                  "cứu và xây dựng các thuật toán có khả năng học hỏi từ dữ liệu kinh nghiệm trong quá khứ để đưa ra quyết "
                  "định hoặc khám phá quy luật trên dữ liệu mới mà không cần lập trình quy tắc tường minh. Đề tài thuộc bài "
                  "toán Học không giám sát (Unsupervised Learning) dạng Phân cụm dữ liệu (Clustering): Gom cụm tập hợp N "
                  "trạm xe đạp dựa trên vector 8 đặc trưng tỷ lệ sử dụng theo khung giờ và ngày trong tuần sao cho các trạm "
                  "trong cùng một cụm có đặc tính nhu cầu tương đồng cao nhất và các trạm khác cụm có sự phân hóa rõ rệt nhất, "
                  "hoàn toàn không cần nhãn giám sát từ trước.")

print("Headings updated successfully!")

# ==============================================================================
# 2. STANDARDIZE ALL TABLES (Captions above + Uniform APA Style)
# ==============================================================================
print("Styling tables and captions...")

tables_meta = [
    {
        "tbl": doc.tables[0],
        "prefix": "Bảng 0.1",
        "title": "Danh mục các từ viết tắt sử dụng trong đề tài",
        "widths": [0.6, 1.3, 2.0, 2.6],
        "aligns": [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]
    },
    {
        "tbl": doc.tables[1],
        "prefix": "Bảng 3.1",
        "title": "So sánh mô hình cơ sở Volume Baseline với mô hình Profile K-Means (K = 3)",
        "widths": [1.8, 2.0, 1.3, 1.4],
        "aligns": [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
    },
    {
        "tbl": doc.tables[2],
        "prefix": "Bảng 3.2",
        "title": "Khảo sát các chỉ số đánh giá chất lượng phân cụm theo số cụm K từ 2 đến 8",
        "widths": [1.1, 1.3, 1.3, 1.4, 1.4],
        "aligns": [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
    },
    {
        "tbl": doc.tables[3],
        "prefix": "Bảng 3.3",
        "title": "Bảng giá trị trung vị (Median Profile) và đặc trưng hành vi của 3 cụm trạm",
        "widths": [2.3, 1.4, 1.4, 1.4],
        "aligns": [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
    },
    {
        "tbl": doc.tables[4],
        "prefix": "Bảng 3.4",
        "title": "Kết quả đánh giá mô hình phân cụm trên tập Test độc lập (N = 335 trạm)",
        "widths": [1.8, 1.5, 1.4, 1.8],
        "aligns": [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
    }
]

for item in tables_meta:
    tbl = item["tbl"]
    add_table_caption_above(tbl, item["prefix"], item["title"])
    format_academic_table(tbl, item["widths"], item["aligns"])

print("All 5 tables formatted identically!")

# ==============================================================================
# 3. STANDARDIZE FIGURES & INSERT CHAPTER 4 IMAGES
# ==============================================================================
print("Standardizing figures and inserting Chapter 4 images...")

# Locate Chapter 4 start index
ch4_start_idx = 0
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("CHƯƠNG 4:"):
        ch4_start_idx = idx
        break

print(f"Chapter 4 begins at paragraph index: {ch4_start_idx}")

# Insert 4 Chapter 4 figures BEFORE their captions in Chapter 4
ch4_fig_inserts = [
    ("Hình 4.1:", os.path.join(FIG_DIR, "fig4_1_architecture.png")),
    ("Hình 4.2:", os.path.join(FIG_DIR, "fig4_2_screen1_intro.png")),
    ("Hình 4.3:", os.path.join(FIG_DIR, "fig4_3_screen2_map.png")),
    ("Hình 4.4:", os.path.join(FIG_DIR, "fig4_4_screen3_dashboard.png"))
]

for caption_start, img_path in ch4_fig_inserts:
    for idx in range(ch4_start_idx, len(doc.paragraphs)):
        p = doc.paragraphs[idx]
        if p.text.strip().startswith(caption_start):
            p_new = doc.add_paragraph()
            p_new.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pPr = p_new._p.get_or_add_pPr()
            sp = OxmlElement('w:spacing')
            sp.set(qn('w:before'), '160') # 8pt
            sp.set(qn('w:after'), '40')   # 2pt
            sp.set(qn('w:line'), '240')
            sp.set(qn('w:lineRule'), 'auto')
            pPr.append(sp)
            pPr.append(OxmlElement('w:keepNext'))
            
            run = p_new.add_run()
            run.add_picture(img_path, width=Inches(5.5))
            
            # Move p_new._p right before caption paragraph p._p
            p._p.addprevious(p_new._p)
            print(f"Successfully inserted image {os.path.basename(img_path)} before caption: {p.text[:40]}")
            break

# Standardize existing image paragraphs in Chapter 3
for p in doc.paragraphs:
    blips = p._p.xpath('.//a:blip')
    if blips and not any(r.text for r in p.runs):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p._p.get_or_add_pPr()
        for sp in pPr.xpath('w:spacing'):
            pPr.remove(sp)
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '160') # 8pt
        sp.set(qn('w:after'), '40')   # 2pt
        sp.set(qn('w:line'), '240')
        sp.set(qn('w:lineRule'), 'auto')
        pPr.append(sp)
        for kwn in pPr.xpath('w:keepNext'):
            pPr.remove(kwn)
        pPr.append(OxmlElement('w:keepNext'))

# Standardize figure captions in Chapter 3 and Chapter 4 (after front matter)
fig_captions_dict = {
    "Hình 3.1:": ("Hình 3.1:", "Thí nghiệm 1 - So sánh phân cụm Volume Baseline vs Hồ sơ Tỷ lệ Profile (Tham chiếu K = 3)"),
    "Hình 3.2:": ("Hình 3.2:", "Thí nghiệm 2 - Đồ thị kép đường cong Elbow và Silhouette Score từ K = 2 đến K = 8"),
    "Hình 3.3:": ("Hình 3.3:", "Thí nghiệm 3 - Ma trận chuyển dịch trạm giữa Mùa xuân và Mùa đông trên 2.129 trạm"),
    "Hình 3.4:": ("Hình 3.4:", "Thí nghiệm 4 - Biểu đồ thanh so sánh giá trị trung vị các đặc trưng giữa 3 cụm (Cluster Medians)"),
    "Hình 4.1:": ("Hình 4.1:", "Sơ đồ kiến trúc Serving phân tách Offline Training và Online Serving"),
    "Hình 4.2:": ("Hình 4.2:", "Giao diện Màn hình 1 - Tổng quan đề tài, phạm vi dữ liệu và giải thích 3 cụm nghiên cứu"),
    "Hình 4.3:": ("Hình 4.3:", "Giao diện Màn hình 2 - Thao tác tương tác Bản đồ trạm Leaflet và Hồ sơ trạm"),
    "Hình 4.4:": ("Hình 4.4:", "Giao diện Màn hình 3 - Dashboard trực quan hóa 4 thí nghiệm và Model Card")
}

for idx, p in enumerate(doc.paragraphs):
    if idx > 90: # Only in body chapters
        txt = p.text.strip()
        for prefix, (cap_pref, cap_desc) in fig_captions_dict.items():
            if txt.startswith(prefix):
                set_caption_text(p, cap_pref, cap_desc)
                break

print("All 8 figures and captions standardized!")

# ==============================================================================
# 4. UPDATE FRONT MATTER (Danh mục hình vẽ & Danh mục bảng biểu)
# ==============================================================================
print("Updating front matter lists...")

for p in doc.paragraphs:
    if p.text.strip() in ["MỤC LỤC HÌNH ẢNH", "DANH MỤC HÌNH VẼ"]:
        p.style = "Normal"
        clean_paragraph_indent(p)
        set_paragraph_spacing(p, before_pt=12, after_pt=6, keep_with_next=True)
        p.text = ""
        r = p.add_run("DANH MỤC HÌNH VẼ")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        break

new_fig_list = [
    "Hình 3.1: Thí nghiệm 1 - So sánh phân cụm Volume Baseline vs Hồ sơ Tỷ lệ Profile",
    "Hình 3.2: Thí nghiệm 2 - Đồ thị kép đường cong Elbow và Silhouette Score từ K = 2 đến K = 8",
    "Hình 3.3: Thí nghiệm 3 - Ma trận chuyển dịch trạm giữa Mùa xuân và Mùa đông trên 2.129 trạm",
    "Hình 3.4: Thí nghiệm 4 - Biểu đồ thanh so sánh giá trị trung vị các đặc trưng giữa 3 cụm",
    "Hình 4.1: Sơ đồ kiến trúc Serving phân tách Offline Training và Online Serving",
    "Hình 4.2: Giao diện Màn hình 1 - Tổng quan đề tài, phạm vi dữ liệu và giải thích 3 cụm nghiên cứu",
    "Hình 4.3: Giao diện Màn hình 2 - Thao tác tương tác Bản đồ trạm Leaflet và Hồ sơ trạm",
    "Hình 4.4: Giao diện Màn hình 3 - Dashboard trực quan hóa 4 thí nghiệm và Model Card"
]

new_tbl_list = [
    "Bảng 0.1: Danh mục các từ viết tắt sử dụng trong đề tài",
    "Bảng 3.1: So sánh mô hình cơ sở Volume Baseline với mô hình Profile K-Means (K = 3)",
    "Bảng 3.2: Khảo sát các chỉ số đánh giá chất lượng phân cụm theo số cụm K từ 2 đến 8",
    "Bảng 3.3: Bảng giá trị trung vị (Median Profile) và đặc trưng hành vi của 3 cụm trạm",
    "Bảng 3.4: Kết quả đánh giá mô hình phân cụm trên tập Test độc lập (N = 335 trạm)"
]

fig_item_paragraphs = []
collecting = False
for p in doc.paragraphs:
    if p.text.strip() == "DANH MỤC HÌNH VẼ":
        collecting = True
        continue
    if collecting:
        if p.text.strip() == "DANH MỤC TỪ VIẾT TẮT":
            break
        fig_item_paragraphs.append(p)

for idx, fig_title in enumerate(new_fig_list):
    if idx < len(fig_item_paragraphs):
        p = fig_item_paragraphs[idx]
        p.text = ""
        p.style = "table of figures"
        clean_paragraph_indent(p)
        set_paragraph_spacing(p, before_pt=2, after_pt=2, keep_with_next=False)
        r = p.add_run(fig_title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

cur_pos = len(new_fig_list)
if cur_pos < len(fig_item_paragraphs):
    p = fig_item_paragraphs[cur_pos]
    p.text = ""
    p.style = "Normal"
    clean_paragraph_indent(p)
    set_paragraph_spacing(p, before_pt=14, after_pt=6, keep_with_next=True)
    r = p.add_run("DANH MỤC BẢNG BIỂU")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    cur_pos += 1

for tbl_title in new_tbl_list:
    if cur_pos < len(fig_item_paragraphs):
        p = fig_item_paragraphs[cur_pos]
        p.text = ""
        p.style = "table of figures"
        clean_paragraph_indent(p)
        set_paragraph_spacing(p, before_pt=2, after_pt=2, keep_with_next=False)
        r = p.add_run(tbl_title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        cur_pos += 1

# Clear any extra placeholder paragraphs
while cur_pos < len(fig_item_paragraphs):
    p = fig_item_paragraphs[cur_pos]
    p.text = ""
    set_paragraph_spacing(p, before_pt=0, after_pt=0, keep_with_next=False)
    cur_pos += 1

# Ensure clean page break before DANH MỤC TỪ VIẾT TẮT
if fig_item_paragraphs:
    last_p = fig_item_paragraphs[-1]
    last_pPr = last_p._p.get_or_add_pPr()
    for br in last_pPr.xpath('.//w:br'):
        last_pPr.remove(br)
    p_run = last_p.add_run()
    p_run.add_break(WD_BREAK.PAGE)

# Format DANH MỤC TỪ VIẾT TẮT heading
for p in doc.paragraphs:
    if p.text.strip() == "DANH MỤC TỪ VIẾT TẮT":
        p.style = "Normal"
        clean_paragraph_indent(p)
        set_paragraph_spacing(p, before_pt=12, after_pt=6, keep_with_next=True)
        p.text = ""
        r = p.add_run("DANH MỤC TỪ VIẾT TẮT")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        break

# Save to both target locations
print("Saving standardized document...")
doc.save(OUT_DOCX_PRIMARY)
doc.save(OUT_DOCX_BACKUP)
print(f"SUCCESS! File saved to {OUT_DOCX_PRIMARY} and {OUT_DOCX_BACKUP}")
